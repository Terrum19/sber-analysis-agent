"""Экспорт исследований и дашбордов в Excel (openpyxl).

Поддерживает:
1. Одиночные исследования/таблицы: карточка запроса, фирменная палитра Сбера,
   автоширина колонок, числовые форматы, автофильтр, фиксация шапки,
   детерминированный подбор диаграмм по Vega-Lite encoding или данным.
2. Дашборды: титульный лист «Сводка» с крупными KPI-карточками, оглавлением
   и ссылками, а также отдельные листы под каждую плитку со снимком данных и графиком.
3. Полная защита от Formula Injection (CWE-1236) на всех строковых поверхностях.
4. Безопасная работа с timezone-aware датами, NaN, Infinity и усечением массивов.
"""

from __future__ import annotations

import logging
import re
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.chart import AreaChart, BarChart, LineChart, PieChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.core.dashboards.models import Dashboard
from app.core.data.duckdb_manager import DuckDBManager
from app.core.export.charts import is_table_spec
from app.core.export.common import (
    extract_kpi_value,
    extract_table_columns,
    format_kpi_number,
    resolve_tile_records,
    sanitize_excel_text,
    sanitize_excel_value,
)

logger = logging.getLogger(__name__)

# Фирменная палитра Сбера
SBER_GREEN = "088434"
SBER_LIGHT_GREEN = "21A038"
SBER_BG_MINT = "F0F7F1"
ZEBRA_FILL_COLOR = "F7FAF7"
BORDER_COLOR = "D9D9D9"
KPI_BORDER_COLOR = "D0DED1"
TEXT_MUTED = "555555"

FONT_NAME = "Calibri"

HEADER_FONT = Font(name=FONT_NAME, size=11, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor=SBER_GREEN)

ZEBRA_FILL = PatternFill("solid", fgColor=ZEBRA_FILL_COLOR)
NO_FILL = PatternFill(fill_type=None)

THIN_BORDER_SIDE = Side(style="thin", color=BORDER_COLOR)
CELL_BORDER = Border(
    left=THIN_BORDER_SIDE,
    right=THIN_BORDER_SIDE,
    top=THIN_BORDER_SIDE,
    bottom=THIN_BORDER_SIDE,
)

KPI_BORDER_SIDE = Side(style="thin", color=KPI_BORDER_COLOR)
KPI_BORDER = Border(
    left=KPI_BORDER_SIDE,
    right=KPI_BORDER_SIDE,
    top=KPI_BORDER_SIDE,
    bottom=KPI_BORDER_SIDE,
)

MAX_EXPORT_ROWS = 5000
MAX_EXCEL_COLUMNS = 50
MAX_CHART_CATEGORIES = 50

# Алиасы для обратной совместимости
_sanitize_value = sanitize_excel_value
_format_kpi_number = format_kpi_number
_extract_kpi_value = extract_kpi_value


def _safe_sheet_title(title: str, existing_titles: set[str]) -> str:
    """Очистить заголовок для имени листа Excel (до 31 символа, без спецсимволов)."""
    clean = re.sub(r"[\\/*?:\[\]]", " ", title).strip()
    if not clean:
        clean = "Лист"
    base = clean[:26].strip()
    candidate = base
    counter = 1
    while candidate.lower() in {t.lower() for t in existing_titles}:
        counter += 1
        candidate = f"{base[:23]}_{counter}"
    existing_titles.add(candidate)
    return candidate


def _is_number(v: Any) -> bool:
    return isinstance(v, (int, float, Decimal)) and not isinstance(v, bool)


def _detect_column_roles(
    records: list[dict], columns: list[str], spec: dict | None = None
) -> tuple[int, list[int]]:
    """Определить 1-based индекс категориальной колонки (X) и числовых метрик (Y).

    Приоритет:
    1. Спецификация Vega-Lite (spec.encoding) для x/y или theta/color;
    2. Умная эвристика по данным выборки с парсингом числовых строк и поддержкой
       всех числовых колонок.
    """
    if spec and isinstance(spec, dict):
        encoding = spec.get("encoding")
        if isinstance(encoding, dict):
            # Круговая диаграмма (pie / donut)
            theta_enc = encoding.get("theta")
            color_enc = encoding.get("color")
            if isinstance(theta_enc, dict) and isinstance(color_enc, dict):
                t_field = theta_enc.get("field")
                c_field = color_enc.get("field")
                if c_field in columns and t_field in columns:
                    return columns.index(c_field) + 1, [columns.index(t_field) + 1]

            # Декартовы диаграммы (bar, line, area, scatter)
            x_enc = encoding.get("x")
            y_enc = encoding.get("y")
            if isinstance(x_enc, dict) and isinstance(y_enc, dict):
                x_field = x_enc.get("field")
                y_field = y_enc.get("field")
                if x_field in columns and y_field in columns:
                    x_idx = columns.index(x_field) + 1
                    y_idx = columns.index(y_field) + 1
                    x_type = (x_enc.get("type") or "").lower()
                    y_type = (y_enc.get("type") or "").lower()

                    # Горизонтальный бар: Y - категория, X - метрика
                    if y_type in ("nominal", "ordinal") and x_type in ("quantitative",):
                        return y_idx, [x_idx]
                    return x_idx, [y_idx]

    # Эвристический подбор по выборке
    sample = records[:50]
    numeric_cols: list[int] = []
    text_or_date_cols: list[int] = []

    def is_val_numeric(v: Any) -> bool:
        if isinstance(v, (int, float, Decimal)) and not isinstance(v, bool):
            return True
        if isinstance(v, str):
            clean = v.strip().replace(" ", "").replace(",", ".").replace("%", "")
            try:
                float(clean)
                return True
            except ValueError:
                return False
        return False

    for idx, col in enumerate(columns, start=1):
        vals = [r.get(col) for r in sample if r.get(col) is not None]
        if not vals:
            continue
        num_count = sum(1 for v in vals if is_val_numeric(v))
        if num_count / len(vals) >= 0.7:
            numeric_cols.append(idx)
        else:
            text_or_date_cols.append(idx)

    if text_or_date_cols:
        cat_col = text_or_date_cols[0]
        metric_cols = [c for c in numeric_cols if c != cat_col]
    elif numeric_cols:
        cat_col = numeric_cols[0]
        metric_cols = numeric_cols[1:]
    else:
        cat_col = 1
        metric_cols = []

    # Исключаем сугубо технические ID, если есть другие метрики
    if len(metric_cols) > 1:
        non_id = [
            c
            for c in metric_cols
            if columns[c - 1].lower()
            not in ("id", "user_id", "order_id", "session_id", "код", "номер", "№", "index")
            and not columns[c - 1].lower().endswith("_id")
        ]
        if non_id:
            metric_cols = non_id

    # Если колонка одна, строим диаграмму по ней
    if not metric_cols and numeric_cols:
        cat_col = numeric_cols[0]
        metric_cols = [numeric_cols[0]]

    return cat_col, metric_cols[:5]


def _detect_chart_type(spec: dict | None, category_col_name: str) -> str:
    """Определить тип диаграммы openpyxl ('bar', 'line', 'pie', 'area')."""
    if spec:
        usermeta = spec.get("usermeta") or {}
        explicit_type = (
            usermeta.get("chart_type") or usermeta.get("requested_type") or spec.get("chart_type")
        )
        if isinstance(explicit_type, str):
            t = explicit_type.lower()
            if t in ("line", "temporal"):
                return "line"
            if t in ("pie", "arc", "donut"):
                return "pie"
            if t in ("area",):
                return "area"
            if t in ("bar", "column", "histogram"):
                return "bar"
            if t in ("scatter", "point"):
                return "line"

        mark = spec.get("mark")
        mark_type = mark.get("type") if isinstance(mark, dict) else mark
        if isinstance(mark_type, str):
            mt = mark_type.lower()
            if mt in ("line", "trail"):
                return "line"
            if mt in ("arc",):
                return "pie"
            if mt in ("area",):
                return "area"
            if mt in ("bar",):
                return "bar"
            if mt in ("point", "circle", "square", "tick"):
                return "line"

    cat_lower = category_col_name.lower()
    time_keywords = (
        "date",
        "дата",
        "месяц",
        "год",
        "квартал",
        "day",
        "день",
        "период",
        "time",
        "year",
        "month",
    )
    if any(k in cat_lower for k in time_keywords):
        return "line"
    return "bar"


def _build_chart(
    ws: Worksheet,
    cat_col: int,
    num_cols: list[int],
    header_row: int,
    last_row: int,
    chart_type: str = "bar",
    title: str | None = None,
) -> Any | None:
    """Создать и настроить объект диаграммы openpyxl."""
    if not num_cols or last_row <= header_row:
        return None

    try:
        if chart_type == "line":
            chart: Any = LineChart()
            chart.style = 13
        elif chart_type == "pie":
            chart = PieChart()
            chart.style = 10
        elif chart_type == "area":
            chart = AreaChart()
            chart.style = 11
        else:
            chart = BarChart()
            chart.type = "col"
            chart.grouping = "clustered"
            chart.style = 10

        if title:
            clean_title = sanitize_excel_text(title[:40])
            chart.title = clean_title.lstrip("'")

        # Ограничиваем диапазон точек для графика (до 50 точек для читаемости)
        chart_max_row = min(last_row, header_row + MAX_CHART_CATEGORIES)
        cats = Reference(ws, min_col=cat_col, min_row=header_row + 1, max_row=chart_max_row)
        chart.set_categories(cats)

        # Для PieChart допускается только 1 серия данных
        cols_to_add = [num_cols[0]] if chart_type == "pie" else num_cols
        for col_idx in cols_to_add:
            data_ref = Reference(ws, min_col=col_idx, min_row=header_row, max_row=chart_max_row)
            chart.add_data(data_ref, titles_from_data=True)

        chart.width = 16
        chart.height = 10
        return chart
    except Exception as exc:  # noqa: BLE001
        logger.warning("Не удалось создать диаграмму Excel: %s", exc)
        return None


def _write_data_table(
    ws: Worksheet,
    records: list[dict],
    columns: list[str],
    header_row: int,
) -> int:
    """Записать стилизованную таблицу данных на лист Excel и вернуть номер последней строки."""
    for col_idx, col_name in enumerate(columns, start=1):
        clean_col = sanitize_excel_text(col_name)
        c = ws.cell(row=header_row, column=col_idx, value=clean_col)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = CELL_BORDER

    cur_row = header_row + 1
    for r_idx, r in enumerate(records):
        is_even = r_idx % 2 == 1
        row_fill = ZEBRA_FILL if is_even else NO_FILL

        for col_idx, col_name in enumerate(columns, start=1):
            raw_val = r.get(col_name) if isinstance(r, dict) else None
            val = sanitize_excel_value(raw_val)
            c = ws.cell(row=cur_row, column=col_idx, value=val)
            c.fill = row_fill
            c.border = CELL_BORDER
            c.font = Font(name=FONT_NAME, size=11)

            if _is_number(val):
                c.alignment = Alignment(horizontal="right", vertical="center")
                if isinstance(val, (float, Decimal)):
                    c.number_format = "#,##0.00"
                else:
                    c.number_format = "#,##0"
            elif isinstance(val, (datetime, date)):
                c.alignment = Alignment(horizontal="center", vertical="center")
                c.number_format = "yyyy-mm-dd"
            elif isinstance(val, bool):
                c.alignment = Alignment(horizontal="center", vertical="center")
            else:
                c.alignment = Alignment(
                    horizontal="left",
                    vertical="center",
                    wrap_text=len(str(val)) > 40,
                )

        cur_row += 1

    last_data_row = cur_row - 1
    ws.freeze_panes = f"A{header_row + 1}"
    last_col_letter = get_column_letter(len(columns))
    ws.auto_filter.ref = f"A{header_row}:{last_col_letter}{last_data_row}"
    return last_data_row


def _auto_adjust_column_widths(
    ws: Worksheet, min_col: int = 1, max_col: int | None = None, min_row: int = 1
) -> None:
    """Автоматический расчёт ширины столбцов по длине содержимого с отступом."""
    if max_col is None:
        max_col = ws.max_column

    for col_idx in range(min_col, max_col + 1):
        max_len = 0
        col_letter = get_column_letter(col_idx)
        for row_idx in range(min_row, ws.max_row + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            if cell.value is not None:
                val_str = str(cell.value)
                for line in val_str.split("\n"):
                    max_len = max(max_len, len(line))
        width = min(max(int(max_len * 1.15) + 4, 11), 60)
        ws.column_dimensions[col_letter].width = width


class ExcelExporter:
    """Выгрузка результатов исследований и дашбордов в формат .xlsx."""

    def export(
        self,
        records: list[dict],
        path: str,
        title: str = "Отчёт",
        question: str | None = None,
        headline: str | None = None,
        details: list[str] | None = None,
        spec: dict | None = None,
    ) -> str:
        """Экспорт отдельного результата исследования (таблицы с диаграммой)."""
        wb = openpyxl.Workbook()
        ws = wb.active
        clean_title = sanitize_excel_text(title)
        ws.title = _safe_sheet_title(clean_title, set())

        cur_row = 1
        total_rows = len(records)
        display_records = records[:MAX_EXPORT_ROWS]

        # Информационная карточка отчёта
        if headline or question or details:
            if headline:
                clean_h = sanitize_excel_text(headline)
                cell = ws.cell(row=cur_row, column=1, value=clean_h)
                cell.font = Font(name=FONT_NAME, size=14, bold=True, color=SBER_GREEN)
                cell.alignment = Alignment(wrap_text=True)
                cur_row += 1

            if question:
                clean_q = sanitize_excel_text(f"Вопрос: {question}")
                cell = ws.cell(row=cur_row, column=1, value=clean_q)
                cell.font = Font(name=FONT_NAME, size=11, bold=True, color="333333")
                cell.alignment = Alignment(wrap_text=True)
                cur_row += 1

            if details:
                for d in details:
                    clean_d = sanitize_excel_text(f"• {d}")
                    cell = ws.cell(row=cur_row, column=1, value=clean_d)
                    cell.font = Font(name=FONT_NAME, size=10, italic=True, color="444444")
                    cell.alignment = Alignment(wrap_text=True)
                    cur_row += 1

            meta_parts = [f"Сформировано: {datetime.now().strftime('%d.%m.%Y %H:%M')}"]
            if total_rows > MAX_EXPORT_ROWS:
                meta_parts.append(
                    f"Показано строк: {len(display_records):,} из {total_rows:,} (лимит выборки)"
                )
            else:
                meta_parts.append(f"Записей: {total_rows}")

            cell = ws.cell(row=cur_row, column=1, value=" | ".join(meta_parts))
            cell.font = Font(name=FONT_NAME, size=9, color=TEXT_MUTED)
            cur_row += 2

        if not display_records:
            empty_cell = ws.cell(row=cur_row, column=1, value="Нет данных для отображения")
            empty_cell.font = Font(name=FONT_NAME, size=11, italic=True, color=TEXT_MUTED)
            Path(path).parent.mkdir(parents=True, exist_ok=True)
            wb.save(path)
            return path

        header_row = cur_row
        columns = extract_table_columns(display_records, max_columns=MAX_EXCEL_COLUMNS)

        last_data_row = _write_data_table(
            ws=ws,
            records=display_records,
            columns=columns,
            header_row=header_row,
        )

        # Построение диаграммы (если спецификация не является чистой таблицей)
        if not is_table_spec(spec):
            cat_col, num_cols = _detect_column_roles(display_records, columns, spec=spec)
            cat_name = columns[cat_col - 1] if cat_col <= len(columns) else ""
            chart_type = _detect_chart_type(spec, cat_name)
            chart = _build_chart(
                ws=ws,
                cat_col=cat_col,
                num_cols=num_cols,
                header_row=header_row,
                last_row=last_data_row,
                chart_type=chart_type,
                title=clean_title if clean_title != "Отчёт" else None,
            )

            if chart is not None:
                chart_col = get_column_letter(len(columns) + 2)
                try:
                    ws.add_chart(chart, f"{chart_col}{header_row}")
                except Exception as exc:  # noqa: BLE001
                    logger.warning("Не удалось добавить диаграмму на лист: %s", exc)

        _auto_adjust_column_widths(ws, min_col=1, max_col=len(columns), min_row=header_row)

        Path(path).parent.mkdir(parents=True, exist_ok=True)
        wb.save(path)
        return path

    def export_dashboard(
        self,
        dashboard: Dashboard,
        path: str,
        db: DuckDBManager | None = None,
        session_title: str | None = None,
        source_name: str | None = None,
        database_name: str | None = None,
    ) -> str:
        """Экспорт целостного дашборда в многостраничный .xlsx файл.

        Включает:
        - Лист «Сводка» с крупными KPI-карточками, реестром плиток и гиперссылками.
        - Отдельные листы под каждую плитку с исходными данными и графиком.

        Скрытые плитки (график убран из дашборда, где построен) не выгружаются.
        """
        dashboard = dashboard.model_copy(
            update={"tiles": [tile for tile in dashboard.tiles if not tile.hidden]}
        )
        wb = openpyxl.Workbook()
        summary_ws = wb.active
        summary_ws.title = "Сводка"

        used_sheet_names: set[str] = {"Сводка"}

        # 1. Шапка листа «Сводка»
        clean_dash_title = sanitize_excel_text(dashboard.title)
        cell = summary_ws.cell(row=1, column=1, value=clean_dash_title)
        cell.font = Font(name=FONT_NAME, size=16, bold=True, color=SBER_GREEN)

        formatted_date = datetime.fromtimestamp(dashboard.created_at).strftime("%d.%m.%Y %H:%M")
        source_label = database_name or session_title or source_name
        if not source_label and db is not None:
            for tile in dashboard.tiles:
                if tile.result_table and not tile.result_table.startswith("temp_result_"):
                    try:
                        from app.core.memory.profile_memory import ProfileMemory

                        dname = ProfileMemory(db).load_display_name(tile.result_table)
                        if dname and dname.strip():
                            source_label = dname.strip()
                            break
                    except Exception as exc:  # noqa: BLE001
                        logger.debug("Не удалось загрузить display_name: %s", exc)

        meta_parts = [f"Дата среза: {formatted_date}"]
        if source_label:
            meta_parts.append(f"База данных: {source_label}")
        elif dashboard.session_id:
            meta_parts.append(f"Сессия: {dashboard.session_id}")
        meta_parts.append(f"Всего плиток: {len(dashboard.tiles)}")

        cell_meta = summary_ws.cell(row=2, column=1, value=" | ".join(meta_parts))
        cell_meta.font = Font(name=FONT_NAME, size=10, color=TEXT_MUTED)

        cur_row = 4

        # 2. Блок KPI-карточек
        kpi_tiles = [
            t
            for t in dashboard.tiles
            if t.is_kpi or (isinstance(t.snapshot_data, dict) and "value" in t.snapshot_data)
        ]

        if kpi_tiles:
            kpi_title_cell = summary_ws.cell(
                row=cur_row, column=1, value="Ключевые показатели (KPI)"
            )
            kpi_title_cell.font = Font(name=FONT_NAME, size=12, bold=True, color=SBER_GREEN)
            cur_row += 1

            cards_per_row = 4
            card_col_width = 2

            for idx, tile in enumerate(kpi_tiles):
                grid_r = idx // cards_per_row
                grid_c = idx % cards_per_row

                start_col = 1 + grid_c * card_col_width
                end_col = start_col + card_col_width - 1
                card_top_row = cur_row + grid_r * 4

                val_str, label_str = _extract_kpi_value(tile)

                summary_ws.merge_cells(
                    start_row=card_top_row,
                    start_column=start_col,
                    end_row=card_top_row + 1,
                    end_column=end_col,
                )
                val_cell = summary_ws.cell(row=card_top_row, column=start_col, value=val_str)
                val_cell.font = Font(name=FONT_NAME, size=18, bold=True, color=SBER_GREEN)
                val_cell.alignment = Alignment(horizontal="center", vertical="center")

                summary_ws.merge_cells(
                    start_row=card_top_row + 2,
                    start_column=start_col,
                    end_row=card_top_row + 2,
                    end_column=end_col,
                )
                clean_lbl = sanitize_excel_text(label_str)
                lbl_cell = summary_ws.cell(row=card_top_row + 2, column=start_col, value=clean_lbl)
                lbl_cell.font = Font(name=FONT_NAME, size=9, bold=False, color=TEXT_MUTED)
                lbl_cell.alignment = Alignment(horizontal="center", vertical="center")

                card_fill = PatternFill("solid", fgColor=SBER_BG_MINT)
                for r_idx in range(card_top_row, card_top_row + 3):
                    for c_idx in range(start_col, end_col + 1):
                        c = summary_ws.cell(row=r_idx, column=c_idx)
                        c.fill = card_fill
                        c.border = KPI_BORDER

            rows_taken = ((len(kpi_tiles) - 1) // cards_per_row + 1) * 4
            cur_row += rows_taken + 1

        # 3. Реестр плиток дашборда (Состав)
        cat_header_cell = summary_ws.cell(row=cur_row, column=1, value="Состав дашборда")
        cat_header_cell.font = Font(name=FONT_NAME, size=12, bold=True, color=SBER_GREEN)
        cur_row += 1

        headers = ["№", "Название плитки", "Тип", "Инсайт / Описание", "Переход к данным"]
        for col_idx, h_text in enumerate(headers, start=1):
            c = summary_ws.cell(row=cur_row, column=col_idx, value=h_text)
            c.font = HEADER_FONT
            c.fill = HEADER_FILL
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = CELL_BORDER

        tile_start_row = cur_row
        cur_row += 1

        tile_sheet_map: dict[str, str | None] = {}
        tile_data_map: dict[str, list[dict] | None] = {}
        for tile in dashboard.tiles:
            records = resolve_tile_records(tile, db=db, limit=MAX_EXPORT_ROWS + 1)

            if isinstance(records, list) and records:
                s_name = _safe_sheet_title(tile.title, used_sheet_names)
                tile_sheet_map[tile.id] = s_name
                tile_data_map[tile.id] = records
            else:
                tile_sheet_map[tile.id] = None
                tile_data_map[tile.id] = None

        for idx, tile in enumerate(dashboard.tiles, start=1):
            is_even = idx % 2 == 0
            r_fill = ZEBRA_FILL if is_even else NO_FILL

            type_label = (
                "KPI"
                if tile.is_kpi
                else (
                    tile.spec.get("usermeta", {}).get("chart_type")
                    or tile.spec.get("chart_type")
                    or "График"
                )
            )

            clean_tile_title = sanitize_excel_text(tile.title)
            clean_type = sanitize_excel_text(type_label)
            clean_insight = sanitize_excel_text(tile.insight or "—")

            row_vals = [
                idx,
                clean_tile_title,
                clean_type,
                clean_insight,
            ]

            for col_idx, val in enumerate(row_vals, start=1):
                c = summary_ws.cell(row=cur_row, column=col_idx, value=val)
                c.fill = r_fill
                c.border = CELL_BORDER
                c.font = Font(name=FONT_NAME, size=10)
                if col_idx == 1:
                    c.alignment = Alignment(horizontal="center", vertical="center")
                elif col_idx == 4:
                    c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                else:
                    c.alignment = Alignment(horizontal="left", vertical="center")

            target_sheet = tile_sheet_map.get(tile.id)
            c_link = summary_ws.cell(row=cur_row, column=5)
            c_link.fill = r_fill
            c_link.border = CELL_BORDER
            if target_sheet:
                c_link.value = "Открыть лист →"
                c_link.hyperlink = f"#'{target_sheet}'!A1"
                c_link.font = Font(name=FONT_NAME, size=10, color=SBER_GREEN, underline="single")
                c_link.alignment = Alignment(horizontal="center", vertical="center")
            else:
                c_link.value = "Сводка"
                c_link.font = Font(name=FONT_NAME, size=10, color=TEXT_MUTED)
                c_link.alignment = Alignment(horizontal="center", vertical="center")

            cur_row += 1

        _auto_adjust_column_widths(summary_ws, min_col=1, max_col=5, min_row=tile_start_row)

        # 4. Создание отдельных листов под плитки со снимками данных
        for tile in dashboard.tiles:
            target_sheet = tile_sheet_map.get(tile.id)
            records = tile_data_map.get(tile.id)
            if not target_sheet or not records:
                continue

            ws = wb.create_sheet(title=target_sheet)

            clean_tile_title = sanitize_excel_text(tile.title)
            t_cell = ws.cell(row=1, column=1, value=clean_tile_title)
            t_cell.font = Font(name=FONT_NAME, size=14, bold=True, color=SBER_GREEN)

            back_cell = ws.cell(row=2, column=1, value="← Вернуться в Сводку")
            back_cell.hyperlink = "#'Сводка'!A1"
            back_cell.font = Font(name=FONT_NAME, size=10, color=SBER_GREEN, underline="single")

            tile_cur_row = 3
            if tile.insight:
                clean_ins = sanitize_excel_text(f"Инсайт: {tile.insight}")
                ins_cell = ws.cell(row=tile_cur_row, column=1, value=clean_ins)
                ins_cell.font = Font(name=FONT_NAME, size=10, italic=True, color="444444")
                ins_cell.alignment = Alignment(wrap_text=True)
                tile_cur_row += 1

            if len(records) > MAX_EXPORT_ROWS or (
                isinstance(tile.snapshot_data, list) and len(tile.snapshot_data) > MAX_EXPORT_ROWS
            ):
                total_cnt = (
                    len(tile.snapshot_data)
                    if isinstance(tile.snapshot_data, list)
                    else len(records)
                )
                msg = (
                    f"Показано первых {MAX_EXPORT_ROWS:,} строк из {total_cnt:,}"
                    if total_cnt > MAX_EXPORT_ROWS + 1
                    else f"Показано первых {MAX_EXPORT_ROWS:,} строк (лимит выборки)"
                )
                trunc_cell = ws.cell(row=tile_cur_row, column=1, value=msg)
                trunc_cell.font = Font(name=FONT_NAME, size=9, italic=True, color=TEXT_MUTED)
                tile_cur_row += 1

            tile_cur_row += 1

            t_header_row = tile_cur_row
            columns = extract_table_columns(records, max_columns=MAX_EXCEL_COLUMNS)

            t_last_row = _write_data_table(
                ws=ws,
                records=records[:MAX_EXPORT_ROWS],
                columns=columns,
                header_row=t_header_row,
            )

            # График для плитки
            if not is_table_spec(tile.spec):
                cat_col, num_cols = _detect_column_roles(records, columns, spec=tile.spec)
                cat_name = columns[cat_col - 1] if cat_col <= len(columns) else ""
                chart_type = _detect_chart_type(tile.spec, cat_name)
                chart = _build_chart(
                    ws=ws,
                    cat_col=cat_col,
                    num_cols=num_cols,
                    header_row=t_header_row,
                    last_row=t_last_row,
                    chart_type=chart_type,
                    title=clean_tile_title,
                )

                if chart is not None:
                    chart_col = get_column_letter(len(columns) + 2)
                    try:
                        ws.add_chart(chart, f"{chart_col}{t_header_row}")
                    except Exception as exc:  # noqa: BLE001
                        logger.warning("Не удалось добавить диаграмму плитки: %s", exc)

            _auto_adjust_column_widths(ws, min_col=1, max_col=len(columns), min_row=t_header_row)

        Path(path).parent.mkdir(parents=True, exist_ok=True)
        wb.save(path)
        return path
